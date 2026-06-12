# -*- coding: utf-8 -*-
"""Generate rewritten Section 4.2 Study Area and Data Collection as a Word doc."""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# base style
st = doc.styles['Normal']
st.font.name = 'Times New Roman'
st.font.size = Pt(11)

def h(text, level):
    p = doc.add_heading(text, level=level)
    return p

def para(text, italic=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic = italic
    p.paragraph_format.space_after = Pt(6)
    return p

def bullet(runs):
    """runs = list of (text, bold) tuples"""
    p = doc.add_paragraph(style='List Bullet')
    for t, b in runs:
        r = p.add_run(t)
        r.bold = b
    return p

# ---------------- 4.2 ----------------
h('4.2 Study Area and Data Collection', 1)

# 4.2.1
h('4.2.1 Study Area: Beijing Municipality', 2)
para(
    "The study area is Beijing Municipality, which in 2023 had a permanent population of "
    "approximately 21.83 million and covered 16,410 square kilometres (China Statistical "
    "Yearbook, 2024). The municipality comprises 16 districts spanning a dense urban core, "
    "surrounding suburban districts, and extensive peri-urban and mountainous outer districts. "
    "This full administrative extent is adopted as the modelling boundary, so that the model "
    "captures long-distance and cross-district commuting between the periphery and the core."
)
para(
    "Within this extent, the central urban area of Beijing (1,381 km², home to roughly 11.37 million "
    "residents; The People’s Governments of Beijing Municipality, 2025) concentrates the overwhelming "
    "majority of employment and commuting demand, and is the focus of interpretation, while the model "
    "itself spans the full municipality. Embedding activity-rich and activity-sparse areas within a "
    "single, consistent choice framework lets the analysis of job accessibility, socio-spatial "
    "inequality and land-use–transport interaction span the full functional region. Figure 3 shows "
    "the location of Beijing Municipality, its 16 districts, and the central urban area."
)

# 4.2.2
h('4.2.2 Spatial Unit and Choice Set', 2)
para(
    "For modelling, the municipality is discretised into a regular grid of 16,907 cells of "
    "approximately one kilometre. Each cell carries its job count, resident population and "
    "district membership. Because the grid is regular, the large outer districts contain the most "
    "cells (e.g., Miyun 2,295; Huairou 2,201; Yanqing 2,070), whereas the compact, high-density core "
    "districts contain few (e.g., Dongcheng 41; Xicheng 49; Shijingshan 81). Employment, by contrast, "
    "is heavily concentrated in the core: Chaoyang (≈1.14 million jobs) and Haidian (≈0.97 million) "
    "alone exceed two million jobs, against fewer than 80,000 in each of the far outer districts such "
    "as Miyun and Yanqing. The administrative reach of the model is therefore municipality-wide, while "
    "the commuting demand it explains is concentrated in the central urban area."
)
para(
    "Commuting is modelled as a joint destination-and-mode choice conditional on origin: for every "
    "origin cell, the model predicts how its commuting outflow distributes across candidate workplace "
    "cells and travel modes. Because the full 16,907 × 16,907 origin–destination matrix is "
    "computationally infeasible and almost entirely empty, the model operates on an observed "
    "candidate (choice) set: 11,371 origin cells have a non-empty candidate set, each with on average "
    "2,674 reachable destinations, together covering 97.2% of observed commuting flow. This sparse "
    "choice-set formulation makes municipality-wide modelling at ∼1 km resolution tractable."
)

# 4.2.3
h('4.2.3 Data', 2)
para(
    "The model draws on spatial, transport, occupational and individual-mobility datasets, combining "
    "open sources, administrative and census records, mobile-network and smart-card data, and "
    "manually curated transit information (Table 8). Three aspects of scope and provenance are noted "
    "first, because they bear directly on what the model can and cannot identify (see Section 4.X on "
    "identifiability):"
)
bullet([
    ("Commuting flows (target variable). ", True),
    ("Cell-level origin–destination flows for four daily periods (morning peak, evening peak, "
     "midday, night) are derived from China Unicom mobile-signalling data (2023; ≈7.06 million OD "
     "records). Crucially, these flows are mode-aggregated — they record how many people travel "
     "A→B but not by which mode — which motivates the external anchoring strategy described below.", False),
])
bullet([
    ("Travel time and the choice set. ", True),
    ("Cell-pair travel times and a congestion index (Gaode, 2025; ≈30.4 million pairs) provide "
     "observed (congested) and free-flow car times, define the candidate set, and underpin walking "
     "and transit time construction. Transit (bus) times are built from bus stops, routes and "
     "timetables (Bjbus, 2024–2025) calibrated to a smart-card-measured operating speed; the metro "
     "network is not separately routed in this version (a documented limitation that biases pure-bus "
     "times upward).", False),
])
bullet([
    ("Occupation, education and workplace. ", True),
    ("A 1% population-census microdata sample (2015) linking individual occupation, workplace district "
     "and travel mode, together with a street-level industry surface from the 2008 Economic Census "
     "(recalibrated with 2023 area-of-interest data), supplies the occupational and educational "
     "dimensions that mode-aggregated OD data alone cannot identify. These records are used both to "
     "anchor mode shares and commuting-tolerance heterogeneity and to validate occupation-specific "
     "destination sorting (e.g., blue-collar workers toward industrial sub-centres, white-collar "
     "workers toward technology and finance clusters).", False),
])
para(
    "The remaining static basemap (road network, building footprints, land use, POI/AOI, residential "
    "communities and street-view imagery) characterises the built environment and supports "
    "accessibility measures and an income proxy (residential housing prices from Anjuke, 2025). "
    "Independent public-transit smart-card trips (2019) provide an observed transit OD used to anchor "
    "where bus and metro flows are directed. Table 8 lists all datasets, distinguishing those used "
    "directly in model estimation from those collected for context or robustness; the 2018 private "
    "electric-vehicle OD, for example, was collected but is not used in the present model."
)

# ---------------- Table 8 ----------------
para("Table 8. Datasets collected for the study and their role in the model.", italic=True)

rows = [
    ("ID", "Theme", "Name", "Year", "Source", "Used in model"),
    ("1", "Static Basemap", "City road centrelines / single road network", "2023", "Manual digitisation", "Context"),
    ("2", "Static Basemap", "City historical road network", "2014–2024", "OpenStreetMap", "Context (congestion ABM)"),
    ("3", "Static Basemap", "Point of Interest (POI)", "2023", "Gaode", "Context"),
    ("4", "Static Basemap", "Area of Interest (AOI)", "2023", "Gaode", "Yes — occupation job-surface calibration"),
    ("5", "Static Basemap", "Residential community information (housing price)", "2025", "Anjuke", "Yes — income proxy"),
    ("6", "Static Basemap", "Building footprints", "2023", "Gaode", "Context"),
    ("7", "Static Basemap", "Land use", "2023", "Gaode", "Context"),
    ("8", "Static Basemap", "Street-view imagery", "2023", "Gaode / Baidu", "Context"),
    ("9", "Public Transport", "Bus stops", "2024", "Bjbus", "Yes — transit time"),
    ("10", "Public Transport", "Bus routes", "2024", "Bjbus", "Yes — transit time"),
    ("11", "Public Transport", "Bus timetables", "2025", "Bjbus", "Yes — transit time"),
    ("12", "Individual Mobility", "Traffic congestion index", "2025", "Gaode", "Yes — car time (peak)"),
    ("13", "Individual Mobility", "Cell-pair travel time (cost)", "2025", "Gaode", "Yes — choice set + car time"),
    ("14", "Individual Mobility", "Private electric-car OD", "2018", "Beijing Traffic Management Bureau", "Collected, not used"),
    ("15", "Individual Mobility", "Mobile-signalling OD (commuting flows)", "2023", "China Unicom", "Yes — target variable"),
    ("16", "Individual Mobility", "Public-transit smart-card OD", "2019", "Beijing transit (Yikatong)", "Yes — transit destination anchor"),
    ("17", "Census / Survey", "Population census 1% microdata (occupation × workplace × mode)", "2015", "National Bureau of Statistics", "Yes — mode-share & education anchors; occupation validation"),
    ("18", "Census / Survey", "Economic census street-level industry surface", "2008", "National Bureau of Statistics", "Yes — typed occupation job surface (M_j^o)"),
    ("19", "Census / Survey", "District occupation & education composition", "2015", "Population Census", "Yes — heterogeneity tiers"),
]

table = doc.add_table(rows=len(rows), cols=6)
table.style = 'Light Grid Accent 1'
for i, row in enumerate(rows):
    cells = table.rows[i].cells
    for j, val in enumerate(row):
        cells[j].text = val
        for p in cells[j].paragraphs:
            for r in p.runs:
                r.font.size = Pt(8)
                if i == 0:
                    r.font.bold = True

# ---------------- modelling approach (retained, lightly revised) ----------------
doc.add_paragraph()
para(
    "This study adopts an activity-based modelling perspective on commuting, recognising that "
    "individuals make travel decisions as part of a broader sequence of daily activities. The "
    "conceptual foundation draws on Chapin’s (1968, 1971) theory that individuals select among "
    "activity options according to underlying motivations, together with the time-geographic "
    "arguments of Hägerstrand (1970), Cullen & Godson (1975) and Jones et al. (1983), which "
    "emphasise temporal and spatial constraints."
)
para(
    "The modelling practice is further informed by ALBATROSS, a well-established activity-based "
    "microsimulation model (Arentze et al., 2000), and by MATSim, one of the most widely applied "
    "multi-agent transport simulation frameworks. ALBATROSS represents activity patterns as the "
    "outcome of learning-based, heuristic decision rules that individuals develop through repeated "
    "interaction with the built environment and the transport system; alternatives are screened and "
    "evaluated against perceived constraints through a consideration set. This non-compensatory, "
    "consideration-set logic is "
    "carried directly into the present model through its lexicographic screening stage (Section 4.X). "
    "MATSim, in turn, demonstrates how data-driven applications across diverse cities — from "
    "transport-electrification scenarios in Sweden to ride-hailing dynamics in Birmingham, Alabama, "
    "and large-scale mobile-signalling-based activity reconstruction in the United States — can "
    "reproduce observed travel patterns with strong rank correlation to independent benchmarks, "
    "underscoring the robustness and policy relevance of data-driven approaches."
)
para(
    "Building on these foundations, this study proposes an Integrated Machine Learning and Agent-Based "
    "Modelling System for Job Accessibility (JAMAS). In its present form, the estimation core is a "
    "nested logit choice model over a sparse spatial choice set — an origin-conditional structured "
    "softmax classifier — augmented by a dual-branch neural-network residual that captures "
    "spatiotemporal patterns the structural utility does not. Machine learning is used to recover "
    "patterns of preference and constraint from large-scale spatiotemporal data, while the "
    "behaviourally structured, sign-constrained utility and the agent-based simulation layer represent "
    "how individuals make commuting decisions and how those decisions aggregate. By integrating "
    "data-driven learning with behavioural simulation, JAMAS offers a means of analysing job "
    "accessibility from both individual and systemic perspectives, as shown in Figure 4."
)

out = r'D:\GIT\mesa Gsoc\GSoC-learning-space\models\05_beijing_nested_logit_v4\report\4.2_Study_Area_and_Data_Collection.docx'
doc.save(out)
print('saved:', out)
